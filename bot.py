import os
import shutil
from dotenv import load_dotenv

from telegram import Update
from telegram.ext import (
    Application,
    CommandHandler,
    MessageHandler,
    ContextTypes,
    filters,
)

from docx import Document
from docx.shared import Inches

load_dotenv()

BOT_TOKEN = os.getenv("BOT_TOKEN")

sessions = {}

TEMP_DIR = "temp"

os.makedirs(TEMP_DIR, exist_ok=True)


async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "Welcome!\n\nUse:\n/new Your Document Title"
    )


async def new_doc(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id

    if not context.args:
        await update.message.reply_text(
            "Usage:\n/new Pharmacology Notes"
        )
        return

    title = " ".join(context.args)

    user_temp_dir = os.path.join(TEMP_DIR, str(user_id))
    os.makedirs(user_temp_dir, exist_ok=True)

    sessions[user_id] = {
        "title": title,
        "items": [],
        "temp_dir": user_temp_dir,
    }

    await update.message.reply_text(
        f"Document '{title}' created.\n\n"
        "Forward messages to me.\n"
        "Send /done when finished."
    )


async def collect_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id

    if user_id not in sessions:
        return

    session = sessions[user_id]

    text = update.message.text or update.message.caption or ""

    image_path = None

    if update.message.photo:
        largest_photo = update.message.photo[-1]

        file = await context.bot.get_file(
            largest_photo.file_id
        )

        image_name = (
            f"{update.message.message_id}.jpg"
        )

        image_path = os.path.join(
            session["temp_dir"],
            image_name,
        )

        await file.download_to_drive(image_path)

    session["items"].append(
        {
            "text": text,
            "image": image_path,
        }
    )


async def done(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id

    if user_id not in sessions:
        await update.message.reply_text(
            "No active document."
        )
        return

    session = sessions[user_id]

    document = Document()

    document.add_heading(
        session["title"],
        level=1
    )

    for item in session["items"]:

        text = item["text"]
        image = item["image"]

        if text:
            document.add_paragraph(
                text,
                style="List Bullet"
            )

        if image and os.path.exists(image):
            try:
                document.add_picture(
                    image,
                    width=Inches(4.5)
                )
            except Exception as e:
                print(
                    f"Failed to add image: {e}"
                )

    safe_title = "".join(
        c
        for c in session["title"]
        if c.isalnum() or c in (" ", "-", "_")
    ).strip()

    filename = f"{safe_title}.docx"

    document.save(filename)

    with open(filename, "rb") as file:
        await update.message.reply_document(
            document=file,
            filename=filename
        )

    if os.path.exists(filename):
        os.remove(filename)

    if os.path.exists(session["temp_dir"]):
        shutil.rmtree(session["temp_dir"])

    del sessions[user_id]


async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id

    if user_id in sessions:

        temp_dir = sessions[user_id]["temp_dir"]

        if os.path.exists(temp_dir):
            shutil.rmtree(temp_dir)

        del sessions[user_id]

    await update.message.reply_text(
        "Cancelled."
    )


def main():
    if not BOT_TOKEN:
        raise ValueError(
            "BOT_TOKEN not found in environment variables"
        )

    app = (
        Application.builder()
        .token(BOT_TOKEN)
        .build()
    )

    app.add_handler(
        CommandHandler("start", start)
    )

    app.add_handler(
        CommandHandler("new", new_doc)
    )

    app.add_handler(
        CommandHandler("done", done)
    )

    app.add_handler(
        CommandHandler("cancel", cancel)
    )

    app.add_handler(
        MessageHandler(
            (filters.TEXT | filters.PHOTO)
            & ~filters.COMMAND,
            collect_message,
        )
    )

    print("Bot started...")
    app.run_polling()


if __name__ == "__main__":
    main()